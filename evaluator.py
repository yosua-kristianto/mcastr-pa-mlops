from sklearn.metrics import accuracy_score, classification_report, confusion_matrix
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import numpy as np
from datetime import datetime


def evaluation(model, model_name, time_between, feature_test, label_test):
    """
    Evaluate model, print results, and write a styled .xlsx report.
    """

    # Predictions + metrics
    predicted = model.predict(feature_test)
    accuracy = accuracy_score(label_test, predicted)
    conf_matrix = confusion_matrix(label_test, predicted)

    class_report_str = classification_report(label_test, predicted)  # pretty print
    class_report_dict = classification_report(label_test, predicted, output_dict=True)  # for DataFrame

    # Print to console
    hours = time_between // 3600
    minutes = (time_between % 3600) // 60
    seconds = time_between % 60

    print("\n=== Evaluation Results ===")
    print(f"Accuracy: {accuracy:.4f}")
    print("\nConfusion Matrix:\n", conf_matrix)
    print("\nClassification Report:\n", class_report_str)

    # === Excel Export ===
    timestamp = datetime.now()
    excel_filename = f"report-{model_name}-{timestamp.strftime('%Y%m%d-%H%M%S')}.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "Evaluation Report"

    # --- Styles ---
    title_font = Font(size=16, bold=True, color="FFFFFF")
    header_font = Font(size=12, bold=True, color="FFFFFF")
    bold_font = Font(bold=True)
    center_align = Alignment(horizontal="center", vertical="center")
    wrap = Alignment(wrap_text=True)
    fill_title = PatternFill("solid", fgColor="4F81BD")   # blue
    fill_header = PatternFill("solid", fgColor="305496")  # darker blue
    alt_fill = PatternFill("solid", fgColor="D9E1F2")     # light gray-blue
    border = Border(
        left=Side(border_style="thin", color="000000"),
        right=Side(border_style="thin", color="000000"),
        top=Side(border_style="thin", color="000000"),
        bottom=Side(border_style="thin", color="000000")
    )

    # --- Title ---
    ws.append([f"{model_name} Evaluation Report"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=6)
    cell = ws["A1"]
    cell.font = title_font
    cell.alignment = center_align
    cell.fill = fill_title

    ws.append([f"Compiled at: {timestamp.strftime('%Y-%m-%d %H:%M:%S')}"])
    ws.append([])

    # Accuracy & Time
    ws.append(["Accuracy", accuracy])
    ws.append(["Training Time", f"{hours}h {minutes}m {seconds}s"])
    ws.append([])

    # --- Confusion Matrix ---
    labels = np.unique(np.concatenate((np.asarray(label_test), np.asarray(predicted)))).tolist()
    ws.append(["Confusion Matrix"])
    ws.append([""] + [str(l) for l in labels])
    header_row_idx = ws.max_row
    for col_idx in range(1, ws.max_column + 1):
        c = ws.cell(row=header_row_idx, column=col_idx)
        c.font = header_font
        c.alignment = center_align
        c.fill = fill_header
        c.border = border

    for i, lab in enumerate(labels):
        row_vals = conf_matrix[i].tolist() if i < conf_matrix.shape[0] else [0] * len(labels)
        row = [str(lab)] + row_vals
        ws.append(row)
        for col_idx, val in enumerate(row, 1):
            c = ws.cell(row=ws.max_row, column=col_idx)
            c.alignment = center_align
            c.border = border
            if ws.max_row % 2 == 0:
                c.fill = alt_fill
    ws.append([])

    # --- Classification Report (table) ---
    ws.append(["Classification Report (table)"])
    df_report = pd.DataFrame(class_report_dict).T.reset_index().rename(columns={'index': 'Class'}).fillna('')

    ws.append(df_report.columns.tolist())
    header_row_idx = ws.max_row
    for col_idx in range(1, len(df_report.columns) + 1):
        c = ws.cell(row=header_row_idx, column=col_idx)
        c.font = header_font
        c.alignment = center_align
        c.fill = fill_header
        c.border = border

    for _, row in df_report.iterrows():
        ws.append([str(x) for x in row.tolist()])
        for col_idx, val in enumerate(row, 1):
            c = ws.cell(row=ws.max_row, column=col_idx)
            c.alignment = center_align
            c.border = border
            if ws.max_row % 2 == 0:
                c.fill = alt_fill

    ws.append([])

    # --- Text Report (raw string for readability) ---
    ws.append(["Classification Report (text)"])
    text_row_idx = ws.max_row + 1
    ws.cell(row=text_row_idx, column=1).value = class_report_str
    ws.cell(row=text_row_idx, column=1).alignment = wrap
    ws.merge_cells(start_row=text_row_idx, start_column=1, end_row=text_row_idx + 6, end_column=6)

    # Auto-fit column widths
    for col in ws.columns:
        max_length = 0
        column = col[0].column
        column_letter = get_column_letter(column)
        for cell in col:
            try:
                if cell.value:
                    val = str(cell.value)
                    line_max = max(len(line) for line in val.splitlines())
                    if line_max > max_length:
                        max_length = line_max
            except:
                pass
        ws.column_dimensions[column_letter].width = (max_length + 2)

    # Save workbook
    wb.save(excel_filename)
    print(f"[INFO] Report saved to {excel_filename}")
